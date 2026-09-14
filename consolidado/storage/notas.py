"""Notas de seguimiento / observaciones de lo que dijo el estudiante."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from consolidado.core.normalizacion import (
    normalizar_id,
    periodo_academico_de_texto,
    periodo_academico_nota,
)
from consolidado.paths import PROJECT_ROOT
from consolidado.storage.db import conexion, inicializar_db
from consolidado.storage.modificaciones import _usuario_log


def _enriquecer(nota: dict[str, Any]) -> dict[str, Any]:
    out = dict(nota)
    out["periodo"] = periodo_academico_de_texto(out.get("creado_en"))
    return out


def listar_notas(
    identificacion: str,
    base: Path | None = None,
    *,
    periodo: str | None = None,
) -> list[dict[str, Any]]:
    base = base or PROJECT_ROOT
    inicializar_db(base)
    ident = normalizar_id(identificacion)
    if not ident:
        return []
    with conexion(base) as conn:
        rows = conn.execute(
            """
            SELECT id, identificacion, nota, creado_en, usuario
            FROM seguimiento_notas
            WHERE identificacion = ?
            ORDER BY creado_en DESC, id DESC
            """,
            (ident,),
        ).fetchall()
    notas = [_enriquecer(dict(r)) for r in rows]
    if periodo:
        return [n for n in notas if n.get("periodo") == periodo]
    return notas


def agregar_nota(
    identificacion: str,
    nota: str,
    *,
    usuario: str | None = None,
    base: Path | None = None,
) -> list[dict[str, Any]]:
    base = base or PROJECT_ROOT
    inicializar_db(base)
    ident = normalizar_id(identificacion)
    texto = (nota or "").strip()
    if not ident or not texto:
        return listar_notas(identificacion, base=base)
    ahora = datetime.now().isoformat(timespec="seconds")
    quien = usuario if usuario is not None else _usuario_log.get()
    with conexion(base) as conn:
        conn.execute(
            """
            INSERT INTO seguimiento_notas (identificacion, nota, creado_en, usuario)
            VALUES (?, ?, ?, ?)
            """,
            (ident, texto, ahora, quien),
        )
    return listar_notas(ident, base=base)


def quitar_nota(nota_id: int, base: Path | None = None) -> None:
    base = base or PROJECT_ROOT
    inicializar_db(base)
    with conexion(base) as conn:
        conn.execute("DELETE FROM seguimiento_notas WHERE id = ?", (int(nota_id),))


def resumen_notas(base: Path | None = None) -> dict[str, dict[str, Any]]:
    """identificacion → {n, ultima, ultima_en, usuario}."""
    base = base or PROJECT_ROOT
    inicializar_db(base)
    with conexion(base) as conn:
        rows = conn.execute(
            """
            SELECT identificacion, nota, creado_en, usuario
            FROM seguimiento_notas
            ORDER BY creado_en DESC, id DESC
            """
        ).fetchall()
    out: dict[str, dict[str, Any]] = {}
    for r in rows:
        ident = normalizar_id(r["identificacion"])
        if not ident:
            continue
        if ident not in out:
            out[ident] = {
                "n": 0,
                "ultima": r["nota"] or "",
                "ultima_en": r["creado_en"],
                "usuario": r["usuario"] or "",
            }
        out[ident]["n"] = int(out[ident]["n"]) + 1
    return out


def listar_todas_notas(base: Path | None = None) -> list[dict[str, Any]]:
    base = base or PROJECT_ROOT
    inicializar_db(base)
    with conexion(base) as conn:
        rows = conn.execute(
            """
            SELECT id, identificacion, nota, creado_en, usuario
            FROM seguimiento_notas
            ORDER BY identificacion, creado_en, id
            """
        ).fetchall()
    return [_enriquecer(dict(r)) for r in rows]


def excel_anotaciones_bytes(
    notas: list[dict[str, Any]] | None = None,
    *,
    estudiantes: pl.DataFrame | None = None,
    base: Path | None = None,
) -> bytes:
    """Excel de anotaciones agrupadas por estudiante, con fecha y periodo."""
    filas = notas if notas is not None else listar_todas_notas(base)
    por_est: dict[str, dict[str, str]] = {}
    if estudiantes is not None and estudiantes.height > 0:
        col_id = "Identificación" if "Identificación" in estudiantes.columns else None
        col_nom = "Nombre y apellidos" if "Nombre y apellidos" in estudiantes.columns else None
        col_prog = "Programa" if "Programa" in estudiantes.columns else None
        if col_id:
            for row in estudiantes.iter_rows(named=True):
                ident = normalizar_id(row.get(col_id))
                if not ident or ident in por_est:
                    continue
                por_est[ident] = {
                    "nombre": str(row.get(col_nom) or "").strip() if col_nom else "",
                    "programa": str(row.get(col_prog) or "").strip() if col_prog else "",
                }

    wb = Workbook()
    ws = wb.active
    ws.title = "Anotaciones"
    cabecera = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="0C6B63")
    cols = [
        "Identificación",
        "Nombre y apellidos",
        "Programa",
        "Fecha",
        "Periodo",
        "Usuario",
        "Nota",
    ]
    ws.append(cols)
    for cell in ws[1]:
        cell.font = cabecera
        cell.fill = fondo
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    ordenadas = sorted(
        filas,
        key=lambda n: (
            str(n.get("identificacion") or ""),
            str(n.get("creado_en") or ""),
            int(n.get("id") or 0),
        ),
    )
    for n in ordenadas:
        ident = normalizar_id(n.get("identificacion")) or str(n.get("identificacion") or "")
        extra = por_est.get(ident, {})
        ws.append(
            [
                ident,
                extra.get("nombre") or "",
                extra.get("programa") or "",
                n.get("creado_en") or "",
                n.get("periodo") or periodo_academico_de_texto(n.get("creado_en")) or "",
                n.get("usuario") or "",
                n.get("nota") or "",
            ]
        )

    ultima = max(len(ordenadas) + 1, 1)
    ult_col = get_column_letter(len(cols))
    if ordenadas:
        tabla = Table(displayName="Anotaciones", ref=f"A1:{ult_col}{ultima}")
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tabla)
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 60
    info = wb.create_sheet("Origen")
    info.append(["Campo", "Valor"])
    for cell in info[1]:
        cell.font = cabecera
        cell.fill = fondo
    info.append(["Semestre de notas (hoy)", periodo_academico_nota()])
    info.append(["Filas", len(ordenadas)])
    info.append(
        [
            "Regla",
            "Dic–mayo = YYYY-1 (diciembre cuenta el año siguiente). Jun–nov = YYYY-2.",
        ]
    )
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 72
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
