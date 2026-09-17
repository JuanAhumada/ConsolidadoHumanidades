"""Excel recortado: una versión del consolidado, columnas elegidas y filtro por carrera."""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from typing import Any

import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from consolidado.config.settings import construir_grupos_encabezado
from consolidado.core.charts import (
    COL_PROGRAMA_GRAFICA,
    _partir_categorias,
    filtrar_df_por_carreras,
    programas_disponibles,
    tabla_items_en_columnas,
)
from consolidado.core.constants import max_materias_en_dataframe

COLUMNAS_BASICAS = (
    "Identificación",
    "Nombre y apellidos",
    "Programa",
)

COL_TIPO_BECA = "Tipo de beca o crédito"
COL_MOTIVO_PRIO = "Motivo Prio."
COL_NIVEL_PRIORIDAD = "Nivel prioridad"
COL_COHORTE_GRAD = "Cohorte de graduación"
COL_PENSUM = "Pensum"

FILTROS_EXTRA = (
    ("beca", COL_TIPO_BECA, "Tipo de beca"),
    ("nivel", COL_NIVEL_PRIORIDAD, "Nivel de prioridad"),
    ("cohorte", COL_COHORTE_GRAD, "Cohorte de graduación"),
    ("pensum", COL_PENSUM, "Pensum"),
)

_RE_NOMBRE_ARCHIVO = re.compile(r"[^A-Za-z0-9._-]+")


def grupos_columnas_version(
    df: pl.DataFrame,
    cfg: dict[str, Any],
) -> list[dict[str, Any]]:
    presentes = [str(c) for c in df.columns]
    vistos: set[str] = set()
    grupos: list[dict[str, Any]] = []
    num_mat = max(max_materias_en_dataframe(df), 1)
    for nombre, cols in construir_grupos_encabezado(cfg, num_mat):
        en_grupo = [c for c in cols if c in presentes and c not in vistos]
        if not en_grupo:
            continue
        grupos.append({"nombre": nombre, "columnas": en_grupo})
        vistos.update(en_grupo)
    resto = [c for c in presentes if c not in vistos]
    if resto:
        grupos.append({"nombre": "Otras", "columnas": resto})
    return grupos


def conteos_por_programa(df: pl.DataFrame) -> list[dict[str, Any]]:
    from collections import Counter

    if COL_PROGRAMA_GRAFICA not in df.columns:
        return []
    cont: Counter[str] = Counter()
    for val in df.get_column(COL_PROGRAMA_GRAFICA).to_list():
        if val is None:
            continue
        for parte in _partir_categorias(str(val)):
            if parte:
                cont[parte] += 1
    return [
        {"nombre": nombre, "n": cont[nombre]}
        for nombre in programas_disponibles(df)
    ]


def conteos_categoria(df: pl.DataFrame, columna: str) -> list[dict[str, Any]]:
    from collections import Counter

    if columna not in df.columns:
        return []
    cont: Counter[str] = Counter()
    for val in df.get_column(columna).to_list():
        if val is None:
            continue
        for parte in _partir_categorias(str(val)):
            if parte:
                cont[parte] += 1
    return [
        {"nombre": nombre, "n": n}
        for nombre, n in sorted(cont.items(), key=lambda p: (-p[1], p[0].casefold()))
    ]


def opciones_filtros(df: pl.DataFrame) -> dict[str, list[dict[str, Any]]]:
    return {
        clave: conteos_categoria(df, col)
        for clave, col, _titulo in FILTROS_EXTRA
    }


def _filtrar_columna_partes(
    df: pl.DataFrame,
    columna: str,
    valores: list[str] | None,
) -> pl.DataFrame:
    elegidos = [str(v).strip() for v in (valores or []) if str(v).strip()]
    if not elegidos or columna not in df.columns:
        return df
    claves = set(elegidos)

    def _coincide(val: Any) -> bool:
        if val is None:
            return False
        return any(parte in claves for parte in _partir_categorias(str(val)))

    return df.filter(pl.col(columna).map_elements(_coincide, return_dtype=pl.Boolean))


def aplicar_filtros_parcializado(
    df: pl.DataFrame,
    *,
    programas: list[str] | None = None,
    filtros: dict[str, list[str]] | None = None,
) -> pl.DataFrame:
    recorte = filtrar_df_por_carreras(df, programas)
    filtros = filtros or {}
    for clave, col, _titulo in FILTROS_EXTRA:
        recorte = _filtrar_columna_partes(recorte, col, filtros.get(clave))
    return recorte


def recortar_dataframe(
    df: pl.DataFrame,
    *,
    columnas: list[str],
    programas: list[str] | None,
    filtros: dict[str, list[str]] | None = None,
) -> pl.DataFrame:
    cols = [str(c).strip() for c in columnas if str(c).strip()]
    if not cols:
        raise ValueError("Elija al menos una columna.")
    faltan = [c for c in cols if c not in df.columns]
    if faltan:
        raise ValueError(f"Esas columnas no están en esta versión: {', '.join(faltan[:8])}.")
    recorte = aplicar_filtros_parcializado(df, programas=programas, filtros=filtros)
    return recorte.select(cols)


def filas_becas_separadas(df: pl.DataFrame) -> tuple[list[str], list[dict[str, Any]]]:
    return tabla_items_en_columnas(df, columna=COL_TIPO_BECA, prefijo="Beca")


def filas_motivos_separados(df: pl.DataFrame) -> tuple[list[str], list[dict[str, Any]]]:
    return tabla_items_en_columnas(df, columna=COL_MOTIVO_PRIO, prefijo="Motivo")


def _valor_excel(val: Any) -> Any:
    if val is None:
        return None
    if isinstance(val, bool):
        return "Sí" if val else "No"
    if isinstance(val, (int, float, datetime, date)):
        return val
    texto = str(val).strip()
    if not texto or texto.lower() in {"none", "null", "nan"}:
        return None
    return texto


def _slug_archivo(texto: str, *, max_len: int = 40) -> str:
    nfkd = unicodedata.normalize("NFKD", texto or "")
    ascii_txt = "".join(c for c in nfkd if not unicodedata.combining(c))
    limpio = _RE_NOMBRE_ARCHIVO.sub("_", ascii_txt).strip("._")
    return (limpio or "parcializado")[:max_len]


def nombre_excel_parcializado(
    *,
    periodo: str | None,
    fecha_version: str | None,
    programas: list[str] | None,
) -> str:
    partes = ["parcializado"]
    if periodo:
        partes.append(_slug_archivo(str(periodo), max_len=16))
    if fecha_version:
        partes.append(_slug_archivo(str(fecha_version), max_len=12))
    elegidas = [p for p in (programas or []) if str(p).strip()]
    if len(elegidas) == 1:
        partes.append(_slug_archivo(elegidas[0], max_len=24))
    elif len(elegidas) > 1:
        partes.append(f"{len(elegidas)}carreras")
    return "_".join(partes) + ".xlsx"


def _escribir_hoja(
    wb: Workbook,
    titulo: str,
    columnas: list[str],
    filas: list[list[Any]],
    *,
    tabla_nombre: str,
    cabecera: Font,
    fondo: PatternFill,
    activa: bool = False,
):
    ws = wb.active if activa else wb.create_sheet(titulo)
    if activa:
        ws.title = titulo
    ws.append(columnas)
    for cell in ws[1]:
        cell.font = cabecera
        cell.fill = fondo
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for fila in filas:
        ws.append(fila)
    ultima = max(len(filas) + 1, 1)
    ult_col = get_column_letter(max(len(columnas), 1))
    if filas and columnas:
        tabla = Table(displayName=tabla_nombre, ref=f"A1:{ult_col}{ultima}")
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tabla)
    else:
        ws.auto_filter.ref = f"A1:{ult_col}{ultima}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    for i, nombre in enumerate(columnas, start=1):
        ancho = min(max(len(str(nombre)) + 2, 12), 42)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    return ws


def excel_parcializado_bytes(
    df: pl.DataFrame,
    *,
    columnas: list[str],
    programas: list[str] | None,
    meta: dict[str, Any] | None = None,
    filtros: dict[str, list[str]] | None = None,
    notas: list[dict[str, Any]] | None = None,
) -> bytes:
    filtrado = aplicar_filtros_parcializado(df, programas=programas, filtros=filtros)
    recorte = recortar_dataframe(
        df, columnas=columnas, programas=programas, filtros=filtros
    )
    wb = Workbook()
    cabecera = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="0C6B63")
    cols = list(recorte.columns)
    _escribir_hoja(
        wb,
        "Parcializado",
        cols,
        [[_valor_excel(fila.get(c)) for c in cols] for fila in recorte.iter_rows(named=True)],
        tabla_nombre="Parcializado",
        cabecera=cabecera,
        fondo=fondo,
        activa=True,
    )

    cols_b, becas = filas_becas_separadas(filtrado)
    if becas:
        _escribir_hoja(
            wb,
            "Becas",
            cols_b,
            [[_valor_excel(f.get(c)) for c in cols_b] for f in becas],
            tabla_nombre="Becas",
            cabecera=cabecera,
            fondo=fondo,
        )

    cols_m, motivos = filas_motivos_separados(filtrado)
    if motivos:
        _escribir_hoja(
            wb,
            "Priorizados",
            cols_m,
            [[_valor_excel(f.get(c)) for c in cols_m] for f in motivos],
            tabla_nombre="Priorizados",
            cabecera=cabecera,
            fondo=fondo,
        )

    if notas is not None:
        ids = set()
        if "Identificación" in filtrado.columns:
            for v in filtrado.get_column("Identificación").to_list():
                if v is None:
                    continue
                texto = str(v).strip()
                if texto:
                    ids.add(texto)
        cols_n = [
            "Identificación",
            "Nombre y apellidos",
            "Programa",
            "Fecha",
            "Periodo",
            "Usuario",
            "Nota",
        ]
        por_est: dict[str, dict[str, Any]] = {}
        if "Identificación" in filtrado.columns:
            for row in filtrado.iter_rows(named=True):
                ident = str(row.get("Identificación") or "").strip()
                if ident and ident not in por_est:
                    por_est[ident] = row
        filas_n = []
        for n in notas:
            ident = str(n.get("identificacion") or "").strip()
            if ids and ident not in ids:
                continue
            extra = por_est.get(ident, {})
            filas_n.append(
                [
                    ident,
                    extra.get("Nombre y apellidos") or "",
                    extra.get("Programa") or "",
                    n.get("creado_en") or "",
                    n.get("periodo") or "",
                    n.get("usuario") or "",
                    n.get("nota") or "",
                ]
            )
        _escribir_hoja(
            wb,
            "Anotaciones",
            cols_n,
            filas_n,
            tabla_nombre="Anotaciones",
            cabecera=cabecera,
            fondo=fondo,
        )

    info = wb.create_sheet("Origen")
    info.append(["Campo", "Valor"])
    for cell in info[1]:
        cell.font = cabecera
        cell.fill = fondo
    meta = meta or {}
    filtros = filtros or {}
    elegidas = [p for p in (programas or []) if str(p).strip()]
    filas_origen = [
        ("Versión", meta.get("id")),
        ("Periodo", meta.get("periodo")),
        ("Fecha del corte", meta.get("fecha_version")),
        ("Carreras", ", ".join(elegidas) if elegidas else "Todas"),
        ("Columnas", recorte.width),
        ("Filas", recorte.height),
        ("Filas de becas", len(becas)),
        ("Filas de motivos", len(motivos)),
    ]
    for clave, _col, titulo in FILTROS_EXTRA:
        vals = [v for v in (filtros.get(clave) or []) if str(v).strip()]
        filas_origen.append((titulo, ", ".join(vals) if vals else "Todos"))
    for campo, valor in filas_origen:
        info.append([campo, valor if valor is not None else "—"])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 48

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
